# Builds mixer_cost.xlsx: full analog mixer cost, US delivery, 2026-10-27.
# Run: uv run --no-project --with openpyxl python cost.py
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

LIVE, EST, OWNED = "live", "estimate", "owned"
DK, TAY, MOU = "DigiKey", "Tayda", "Mouser"

# (section, item, part / note, qty, unit USD, source, status)
ROWS = [
    # --- Semiconductors (counts: 4 mono strips, 2 stereo strips, master, digital, reverb, power)
    ("Chips", "Mic preamp", "THAT1512P08-U (4 strips)", 4, 5.84, MOU, LIVE),
    ("Chips", "Op-amps", "NE5532P (12 strips + 4 master + 8 stereo)", 24, 0.55, TAY, EST),
    ("Chips", "Comparators", "LM393 (4 strips + 4 stereo)", 8, 0.35, TAY, EST),
    ("Chips", "Peak op-amp", "TL072CP (master)", 1, 0.60, TAY, EST),
    ("Chips", "Balanced line drivers", "TI DRV134PA (master)", 2, 6.50, MOU, EST),
    ("Chips", "Headphone amp", "NJM4556AD (master)", 1, 1.52, DK, LIVE),
    ("Chips", "Meter drivers", "LM3915N-1 or NTE1509 (master)", 2, 7.00, "Jameco / Mouser", EST),
    ("Chips", "MIDI opto, reverb LDO + op-amp", "6N138, MCP1700-3302E, MCP6002-I/P", 1, 2.00, DK, EST),
    ("Chips", "IC sockets", "43x DIP-8 + 2x DIP-18", 1, 3.50, TAY, EST),
    # --- Pots, faders, switches
    ("Controls", "Fader, mono strips", "Bourns PTA6043-2015DPA103 60 mm", 4, 2.37, DK, LIVE),
    ("Controls", "Fader, dual (master + 2 stereo)", "Bourns PTA6044-2015DPA103 (DigiKey: 2000 MOQ, buy at Mouser/Utmel)", 3, 2.70, "Mouser / Utmel", EST),
    ("Controls", "Pot A10K single", "Alpha RD901F-40-15R1-A10K (Tayda A-1856)", 14, 1.29, TAY, LIVE),
    ("Controls", "Pot C10K gain", "Alpha RD901F-40-15R1-C10K", 4, 1.29, TAY, EST),
    ("Controls", "Pot B10K centre detent", "Alpha RD901F-40-15R1-B10K-0CD (Tayda A-8401)", 18, 1.40, TAY, EST),
    ("Controls", "Pot dual (phones A10K, stereo EQ B10K)", "Alpha RD902F (Tayda A-6980 / A-6433)", 5, 2.20, TAY, EST),
    ("Controls", "Latching switches", "C&K PVA2 EE H1 1.7N V2 (12 strip + 1 master)", 13, 2.424, DK, LIVE),
    ("Controls", "Switch caps", "C&K PVA push caps", 13, 0.60, DK, EST),
    ("Controls", "Reverb trimmers", "Bourns 3296W 10k x4, 20k x1", 5, 1.50, DK, EST),
    # --- Connectors
    ("Connectors", "Mic combo jacks", "Neutrik NCJ6FA-V", 4, 4.50, DK, EST),
    ("Connectors", "Main XLR out", "Neutrik NC3MAAV", 2, 3.50, DK, EST),
    ("Connectors", "1/4\" jacks, master", "Neutrik NJ3FD-V", 4, 1.80, DK, EST),
    ("Connectors", "1/4\" jacks, stereo line strip", "Neutrik NJ5FD-V (USB-return strip has none)", 2, 1.63, DK, LIVE),
    ("Connectors", "3.5 mm jacks", "Thonkiconn PJ366ST (master + stereo line)", 2, 0.50, "Thonk", EST),
    ("Connectors", "Shrouded 2x5 headers", "Wurth 61201021621", 17, 0.45, MOU, EST),
    ("Connectors", "Ribbon cable + IDC plugs", "2x5 IDC sockets x17 + 1 m 10-way ribbon", 1, 7.00, TAY, EST),
    ("Connectors", "Digital board jacks", "USB-B, USB-C (GCT), USB-A, 2x DIN-5", 1, 5.00, DK, EST),
    ("Connectors", "Rear DC jack + pigtail", "Panel 2.1 mm jack, short 2.1 mm plug lead", 1, 3.00, "AliExpress", EST),
    # --- Passives
    ("Passives", "Resistors", "~320 metal film incl. 8x 0.1% 6.81k", 1, 15.00, TAY, EST),
    ("Passives", "Capacitors", "Films, ceramics, 7 mm low-profile electrolytics", 1, 30.00, "Tayda / DigiKey", EST),
    ("Passives", "LEDs, diodes, DIP switch, headers", "3 mm LEDs, 1N4004/4148/5817, pin headers", 1, 5.00, TAY, EST),
    # --- Power board
    ("Power", "AC-AC adapter", "Jameco ReliaPro ADU150150RV0650, 15 VAC 1.5 A", 1, 10.95, "Jameco", LIVE),
    ("Power", "Big caps", "6800uF 35V snap-in x2, 1000uF 63V x4", 1, 11.00, DK, EST),
    ("Power", "Regulators, heatsinks, fuse, rest", "7812/7912, 2x Fischer SK104, MJE340, fuse + holder, barrel jack, insulator kit", 1, 9.00, DK, EST),
    # --- Modules
    ("Modules", "Reverb module", "Akizuki AE-FV1-DIP, JPY 2,780 at ~150 JPY/USD", 1, 18.50, "Akizuki", LIVE),
    ("Modules", "Proxy + Japan shipping", "ZenMarket fee + EMS/small packet to US", 1, 25.00, "ZenMarket", EST),
    ("Modules", "USB hub module", "FE1.1s (Amazon JESSINIE)", 1, 8.00, "Amazon", EST),
    ("Modules", "USB audio module", "PCM2902 module", 1, 8.00, "AliExpress", EST),
    # --- PCBs (JLCPCB, 2-layer, 5 of each)
    ("PCBs", "Mono strip 30x343 mm", "5 pcs (4 used)", 1, 18.00, "JLCPCB", EST),
    ("PCBs", "Stereo strip 30x343 mm", "5 pcs (2 used)", 1, 18.00, "JLCPCB", EST),
    ("PCBs", "Master 60x343 mm", "5 pcs (1 used)", 1, 30.00, "JLCPCB", EST),
    ("PCBs", "Power, digital, reverb, controller", "4 small boards, 5 pcs each", 1, 10.00, "JLCPCB", EST),
    ("PCBs", "Shipping + import duty", "DHL to US, ~1.5 kg", 1, 45.00, "JLCPCB", EST),
    # --- Case
    ("Case", "Printed top (control plate + jack shelf)", "Omer chose printed over aluminium 2026-10-27; PLA ~0.45 kg, split into bed-size tiles", 1, 0.00, "-", EST),
    ("Case", "Filament", "PLA 2 spools: case ~0.8 kg + printed top ~0.45 kg + spare", 2, 18.00, "Amazon", EST),
    ("Case", "Hardware", "M3 F-F standoffs 10.5/7.5 mm, M3x10 countersunk, M2x8 + 4 mm spacers, seam bolts", 1, 15.00, "Amazon", EST),
    ("Case", "M3 screws + heat-set inserts", "Already bought for the controller", 1, 0.00, "-", OWNED),
    # --- Knobs
    ("Knobs", "Pot knobs", "41 knobs, 6.35 mm round shaft", 41, 0.60, TAY, EST),
    ("Knobs", "Fader caps", "7 caps for PTA 15 mm metal lever", 7, 1.00, "AliExpress", EST),
    ("Knobs", "Controller knobs", "4 fader caps + 4 encoder knobs", 1, 8.00, "AliExpress", EST),
    # --- Controller (MIDI section)
    ("Controller", "ESP32-S3 devkit, 2.8\" TFT, 45 mm faders, EC11s, buttons, LEDs, resistors", "Already bought", 1, 0.00, "-", OWNED),
    ("Controller", "Headers", "Female 2.54 mm strips", 1, 2.00, "AliExpress", EST),
    # --- Shipping from parts shops
    ("Shipping", "DigiKey / Mouser", "Free over the order threshold", 1, 0.00, DK, EST),
    ("Shipping", "Tayda", "To US", 1, 6.00, TAY, EST),
    ("Shipping", "Jameco", "Ground", 1, 9.00, "Jameco", EST),
]

wb = Workbook()
ws = wb.active
ws.title = "Cost"
head = ["Section", "Item", "Part / note", "Qty", "Unit $", "Total $", "Where", "Price"]
ws.append(head)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="333333")
fill = {EST: PatternFill("solid", fgColor="FFF2CC"), OWNED: PatternFill("solid", fgColor="E2EFDA")}
for r in ROWS:
    ws.append(list(r[:5]) + [None] + list(r[5:]))
    i = ws.max_row
    ws.cell(i, 6, f"=D{i}*E{i}")
    if r[6] in fill:
        for c in ws[i]:
            c.fill = fill[r[6]]
last = ws.max_row
for col in "EF":
    for c in ws[col][1:]:
        c.number_format = "$#,##0.00"
for col, w in zip("ABCDEFGH", (12, 34, 62, 6, 9, 10, 16, 10)):
    ws.column_dimensions[col].width = w

# Summary sheet with SUMIF per section
s = wb.create_sheet("Summary", 0)
s.append(["Section", "Total $", "of which estimated $"])
for c in s[1]:
    c.font = Font(bold=True)
sections = list(dict.fromkeys(r[0] for r in ROWS))
for sec in sections:
    i = s.max_row + 1
    s.append([sec, f'=SUMIF(Cost!A2:A{last},A{i},Cost!F2:F{last})',
              f'=SUMIFS(Cost!F2:F{last},Cost!A2:A{last},A{i},Cost!H2:H{last},"{EST}")'])
n = s.max_row
s.append(["TOTAL", f"=SUM(B2:B{n})", f"=SUM(C2:C{n})"])
s.append([])
s.append(["Budget target (standard tier)", 300])
s.append(["Over budget", f"=B{n+1}-B{n+3}"])
for row in s.iter_rows(min_row=2, min_col=2, max_col=3):
    for c in row:
        c.number_format = "$#,##0"
for c in s[n + 1]:
    c.font = Font(bold=True)
s.column_dimensions["A"].width = 30
s.column_dimensions["B"].width = 12
s.column_dimensions["C"].width = 20
s.append([])
s.append(["Yellow rows on the Cost sheet are estimates; white rows are prices read on 2026-10-27; green rows Omer already owns."])
s.append(["Not included: sales tax, solder/tools, spares, the optional 24LC32A EEPROM."])
s.cell(s.max_row - 1, 1).alignment = Alignment(wrap_text=False)

wb.save("mixer_cost.xlsx")

total = sum(r[3] * r[4] for r in ROWS)
est = sum(r[3] * r[4] for r in ROWS if r[6] == EST)
by = {sec: sum(r[3] * r[4] for r in ROWS if r[0] == sec) for sec in sections}
assert abs(sum(by.values()) - total) < 1e-6
for sec, v in by.items():
    print(f"{sec:12s} ${v:7.2f}")
print(f"TOTAL ${total:.2f}  (estimated part ${est:.2f})")
